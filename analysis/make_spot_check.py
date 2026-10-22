"""
Build private/spot_check_all_reviews.xlsx: every review with its full text beside the labels the
article uses, so a person can check any of them by hand.

    python analysis/make_spot_check.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
PRIV = ROOT / "private"
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD, YELLOW = PatternFill("solid", fgColor="2F3B52"), PatternFill("solid", fgColor="FFF2CC")
WIDTH = {"Review text": 80, "Title": 28, "Supporting quote": 34, "Listing": 30, "Reviewer's own tags": 24}


def quotes():
    q = pd.read_csv(PRIV / "verification/full_read_labels_with_quotes.csv", dtype=str, keep_default_na=False)
    fixes = pd.read_csv(ROOT / "data/label_corrections.csv", dtype=str)
    q.loc[q.uid.isin(fixes.uid), "quote"] = "(label corrected; see data/label_corrections.csv)"
    return q.set_index("uid").quote


def sheet(wb, name, df):
    ws = wb.create_sheet(name)
    heads = list(df.columns) + ["Your check", "Your notes"]
    ws.append(heads)
    for row in df.itertuples(index=False):
        ws.append([("" if pd.isna(v) else v) for v in row] + ["", ""])
    for i, h in enumerate(heads, 1):
        c = ws.cell(1, i)
        c.font, c.fill, c.alignment = BOLD, HEAD, Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDTH.get(h, 13)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = FONT, Alignment(wrap_text=True, vertical="top")
        row[-2].fill = row[-1].fill = YELLOW
    col = get_column_letter(len(heads) - 1)
    dv = DataValidation(type="list", formula1='"OK,Wrong: fades,Wrong: smoke,Wrong: gifted,Wrong: scent,Text mismatch,Other"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{ws.max_row}")
    ws.freeze_panes, ws.auto_filter.ref = "C2", f"A1:{get_column_letter(len(heads))}{ws.max_row}"
    return ws.max_row, heads


def main():
    q = quotes()
    yn = lambda x: "Yes" if bool(x) else "No"
    sr = pd.read_csv(PRIV / "sephora_reviews_raw.csv", dtype=str, keep_default_na=False)
    sc = pd.read_csv(ROOT / "data/sephora_reviews_coded.csv", dtype={"review_id": str})
    S = sr[["review_id", "title", "text"]].merge(sc, on="review_id").sort_values(["scent", "date"])
    S = pd.DataFrame({"Review ID": S.review_id, "Scent": S.scent, "Listing": S.product_name,
                      "Date": S.date,
                      "Stars": S.rating, "Got it free": S.gifted.map(yn), "Title": S.title, "Review text": S.text,
                      "Fades / too faint": S.fades_read.astype(int), "Cigarette-type smoke": S.smoke_read.astype(int),
                      "Supporting quote": ("sephora:" + S.review_id).map(q)})
    wr = pd.read_csv(PRIV / "lore_world_reviews_raw.csv", dtype=str, keep_default_na=False)
    wc = pd.read_csv(ROOT / "data/lore_world_reviews_coded.csv", dtype={"review_id": str})
    W = wr[["review_id", "title", "text"]].merge(wc, on="review_id").sort_values(["scent", "date"])
    W = pd.DataFrame({"Review ID": W.review_id, "Scent": W.scent, "Listing": W.listings, "Date": W.date,
                      "Stars": W.rating, "Verified buyer": W.verified.astype(str), "Title": W.title, "Review text": W.text,
                      "Fades / too faint": W.fades_read.astype(int), "Cigarette-type smoke": W.smoke_read.astype(int),
                      "Supporting quote": ("lore_world:" + W.review_id).map(q)})
    fr = pd.read_csv(PRIV / "fragrantica_reviews_raw.csv", dtype=str, keep_default_na=False)
    fc = pd.read_csv(ROOT / "data/fragrantica_reviews_coded.csv", dtype={"review_id": str})
    F = fr[["review_id", "tags", "text"]].merge(fc[["review_id", "scent", "date", "fades_read", "smoke_read"]],
                                                on="review_id").sort_values(["scent", "date"])
    F = pd.DataFrame({"Review ID": F.review_id, "Scent": F.scent, "Date": F.date, "Reviewer's own tags": F.tags,
                      "Review text": F.text, "Fades / too faint": F.fades_read.astype(int),
                      "Cigarette-type smoke": F.smoke_read.astype(int), "Supporting quote": ("fragrantica:" + F.review_id).map(q)})

    wb = Workbook()
    guide = wb.active
    guide.title = "How to use"
    info = {}
    for name, df in [("Sephora", S), ("lore.world", W), ("Fragrantica", F)]:
        n, heads = sheet(wb, name, df)
        info[name] = (n, heads)
    lines = [["Spot-check sheet: every LORE review behind the article"], [""],
             ["Each tab holds every review from one source: its full text and the labels the article uses."],
             ["Pick a value in the yellow 'Your check' column, and add a note if useful. Filter the header row by scent or label."],
             ["Fades / too faint = 1 when the reviewer says the scent fades, is faint, projects weakly or wishes it lasted longer."],
             ["Cigarette-type smoke = 1 when the reviewer compares it to cigarettes, an ashtray, BBQ or burnt rubber."],
             ["Got it free = Sephora's gifted flag, or Sephora's '[I received this product...]' line in the review."],
             ["Example: 'love it, just wish it lasted longer' should read Fades = 1; if you agree, choose OK."],
             ["For the reviews the verification script couldn't confirm, start with private/needs_human_review.xlsx."],
             [""], ["Summary", "Reviews", "Fades = 1", "Smoke = 1", "Checked", "Marked wrong"]]
    for line in lines:
        guide.append(line)
    for name, (n, heads) in info.items():
        ref = f"'{name}'"
        fcol = get_column_letter(heads.index("Fades / too faint") + 1)
        scol = get_column_letter(heads.index("Cigarette-type smoke") + 1)
        kcol = get_column_letter(heads.index("Your check") + 1)
        guide.append([name, f"=COUNTA({ref}!A2:A{n})", f"=SUM({ref}!{fcol}2:{fcol}{n})", f"=SUM({ref}!{scol}2:{scol}{n})",
                      f"=COUNTA({ref}!{kcol}2:{kcol}{n})",
                      f"=COUNTIF({ref}!{kcol}2:{kcol}{n},\"Wrong*\")+COUNTIF({ref}!{kcol}2:{kcol}{n},\"Text mismatch\")"])
    guide.append(["Source: Sephora, lore.world and Fragrantica review pulls of Oct 21, 2026"])
    for row in guide.iter_rows():
        for c in row:
            c.font = FONT
    guide["A1"].font = Font(name="Arial", size=14, bold=True)
    for c in guide[11]:
        c.font = Font(name="Arial", size=10, bold=True)
    guide.column_dimensions["A"].width = 34
    wb.save(PRIV / "spot_check_all_reviews.xlsx")
    print("Wrote private/spot_check_all_reviews.xlsx")


if __name__ == "__main__":
    main()
